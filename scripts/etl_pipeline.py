"""
ETL Pipeline for Cancer Intelligence & Analytics
Course: Data Warehousing and Data Mining (DWM) - Sem V CMPN
Database: SQLite at warehouse/cancer_dw.db
Sources:
  - Data/extracted_temp/Cancer Dataset/final_excel_updated.xlsx (Primary)
  - Data/extracted_temp/Cancer Dataset/Dataset 9.xlsx (Unique auxiliary sheets)
  - Data/extracted_temp/Cancer Dataset/Dataset 6.docx (Screening & Approved Infra)
  - Data/extracted_temp/Cancer Dataset/Dataset 5.docx (State-wise Mortality & Incidence)
"""

import os
import re
import sys
import json
import sqlite3
import logging
import openpyxl
import docx
import pandas as pd
import numpy as np

# Configure Logging
log_file = os.path.join("logs", "etl_pipeline.log")
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        logging.FileHandler(log_file, mode="w", encoding="utf-8"),
        logging.StreamHandler(sys.stdout)
    ]
)

DB_PATH = os.path.join("warehouse", "cancer_dw.db")
TEMP_DB_PATH = os.path.join("warehouse", "cancer_dw_TEMP.db")
BASE_DATA_DIR = os.path.join("Data", "extracted_temp", "Cancer Dataset")

FINAL_EXCEL = os.path.join(BASE_DATA_DIR, "final_excel_updated.xlsx")
DATASET_9 = os.path.join(BASE_DATA_DIR, "Dataset 9.xlsx")
DATASET_6 = os.path.join(BASE_DATA_DIR, "Dataset 6.docx")
DATASET_5 = os.path.join(BASE_DATA_DIR, "Dataset 5.docx")

# State Name Normalization Mapping
STATE_CODE_MAP = {
    "Jammu and Kashmir": ("JK", "North"),
    "Jammu & Kashmir": ("JK", "North"),
    "Ladakh": ("LA", "North"),
    "Himachal Pradesh": ("HP", "North"),
    "Punjab": ("PB", "North"),
    "Chandigarh": ("CH", "North"),
    "Uttaranchal": ("UK", "North"),
    "Uttarakhand": ("UK", "North"),
    "Haryana": ("HR", "North"),
    "Delhi": ("DL", "North"),
    "Rajasthan": ("RJ", "Central"),
    "Uttar Pradesh": ("UP", "Central"),
    "Bihar": ("BR", "East"),
    "Sikkim": ("SK", "North East"),
    "Arunachal Pradesh": ("AR", "North East"),
    "Nagaland": ("NL", "North East"),
    "Manipur": ("MN", "North East"),
    "Mizoram": ("MZ", "North East"),
    "Tripura": ("TR", "North East"),
    "Meghalaya": ("ML", "North East"),
    "Assam": ("AS", "North East"),
    "West Bengal": ("WB", "East"),
    "Jharkhand": ("JH", "East"),
    "Orissa": ("OR", "East"),
    "Odisha": ("OR", "East"),
    "Chhattisgarh": ("CG", "Central"),
    "Madhya Pradesh": ("MP", "Central"),
    "Gujarat": ("GJ", "West"),
    "Daman": ("DD", "West"),
    "Daman & Diu": ("DD", "West"),
    "Dadra and Nagar Haveli": ("DN", "West"),
    "Dadra & Nagar Haveli": ("DN", "West"),
    "Maharashtra": ("MH", "West"),
    "Telangana": ("TG", "South"),
    "Andhra Pradesh": ("AP", "South"),
    "Karnataka": ("KA", "South"),
    "Goa": ("GA", "West"),
    "Lakshadweep": ("LD", "South"),
    "Kerala": ("KL", "South"),
    "Tamil Nadu": ("TN", "South"),
    "Pondicherry": ("PY", "South"),
    "Puducherry": ("PY", "South"),
    "Andaman and Nicobar Islands": ("AN", "South"),
    "Andaman & Nicobar Islands": ("AN", "South")
}

def clean_num(val):
    """Safely cast integer/float strings with commas or non-breaking spaces."""
    if val is None or pd.isna(val):
        return None
    if isinstance(val, (int, float)):
        return val
    s = str(val).strip().replace(",", "").replace("\xa0", "")
    if s == "" or s.lower() in ("none", "null", "na", "n/a", "-", "."):
        return None
    try:
        if "." in s:
            return float(s)
        return int(s)
    except ValueError:
        return None

def init_database(conn):
    """Create all Dimension and Fact tables with proper primary/foreign keys."""
    cursor = conn.cursor()
    logging.info("Initializing SQLite warehouse tables...")

    cursor.executescript("""
    DROP VIEW IF EXISTS vw_star_state_incidence;
    DROP VIEW IF EXISTS vw_star_cancer_sites;
    DROP VIEW IF EXISTS vw_state_mortality_ratios;
    
    DROP TABLE IF EXISTS Fact_State_Mortality_Surveillance;
    DROP TABLE IF EXISTS Fact_Cancer_Incidence;
    DROP TABLE IF EXISTS Fact_PBCR_Surveillance;
    DROP TABLE IF EXISTS Fact_State_Surveillance;
    DROP TABLE IF EXISTS Fact_Screening_Programs;
    DROP TABLE IF EXISTS Fact_Clinical_Stage;
    DROP TABLE IF EXISTS Fact_Cancer_Trends_APC;
    DROP TABLE IF EXISTS Fact_Healthcare_Infrastructure_National;
    DROP TABLE IF EXISTS Fact_Maharashtra_Infrastructure;
    
    DROP TABLE IF EXISTS Dim_Geography;
    DROP TABLE IF EXISTS Dim_Cancer_Site;
    DROP TABLE IF EXISTS Dim_Organ_System;
    DROP TABLE IF EXISTS Dim_Time;
    DROP TABLE IF EXISTS Dim_Measure_Type;
    DROP TABLE IF EXISTS Dim_Registry;

    -- Conformed Dimension: Geography
    CREATE TABLE Dim_Geography (
        geo_id INTEGER PRIMARY KEY AUTOINCREMENT,
        state_name TEXT UNIQUE NOT NULL,
        state_code TEXT NOT NULL,
        region_zone TEXT NOT NULL,
        country TEXT DEFAULT 'India'
    );

    -- Snowflake Sub-Dimension: Organ System
    CREATE TABLE Dim_Organ_System (
        organ_system_id INTEGER PRIMARY KEY AUTOINCREMENT,
        system_name TEXT UNIQUE NOT NULL
    );

    -- Dimension: Cancer Site (Normalized to Organ System)
    CREATE TABLE Dim_Cancer_Site (
        site_id INTEGER PRIMARY KEY AUTOINCREMENT,
        site_name TEXT UNIQUE NOT NULL,
        icd10_code TEXT,
        level_type TEXT, -- Group, Site, Total
        parent_group TEXT,
        is_tobacco_associated INTEGER DEFAULT 0, -- 1=True, 0=False
        organ_system_id INTEGER,
        FOREIGN KEY (organ_system_id) REFERENCES Dim_Organ_System(organ_system_id)
    );

    -- Conformed Dimension: Time
    CREATE TABLE Dim_Time (
        time_id INTEGER PRIMARY KEY,
        year INTEGER NOT NULL,
        period_label TEXT NOT NULL,
        epoch_category TEXT NOT NULL -- Historical, Extrapolated, Benchmark, Projected
    );

    -- Dimension: Measure Type (Distinguishes Observed vs Estimated vs Forecast)
    CREATE TABLE Dim_Measure_Type (
        measure_type_id INTEGER PRIMARY KEY AUTOINCREMENT,
        measure_code TEXT UNIQUE NOT NULL,
        measure_name TEXT NOT NULL,
        methodology_source TEXT NOT NULL,
        is_modeled INTEGER NOT NULL -- 1=Model Forecast, 0=Empirical/Extrapolation
    );

    -- Dimension: Registry
    CREATE TABLE Dim_Registry (
        registry_id INTEGER PRIMARY KEY AUTOINCREMENT,
        registry_name TEXT UNIQUE NOT NULL,
        state_id INTEGER,
        region TEXT,
        urban_percent REAL,
        rural_percent REAL,
        established_year INTEGER,
        FOREIGN KEY (state_id) REFERENCES Dim_Geography(geo_id)
    );

    -- Fact 1: State Annual Incidence & Projections (37 States x Multi-Year)
    CREATE TABLE Fact_State_Surveillance (
        fact_state_id INTEGER PRIMARY KEY AUTOINCREMENT,
        geo_id INTEGER NOT NULL,
        time_id INTEGER NOT NULL,
        measure_type_id INTEGER NOT NULL,
        incident_cases REAL NOT NULL,
        cagr_observed REAL,
        cagr_projected REAL,
        national_share_pct REAL,
        FOREIGN KEY (geo_id) REFERENCES Dim_Geography(geo_id),
        FOREIGN KEY (time_id) REFERENCES Dim_Time(time_id),
        FOREIGN KEY (measure_type_id) REFERENCES Dim_Measure_Type(measure_type_id)
    );

    -- Fact 2: National Site-wise Burden (Site x Year x Gender)
    CREATE TABLE Fact_Cancer_Incidence (
        fact_id INTEGER PRIMARY KEY AUTOINCREMENT,
        site_id INTEGER NOT NULL,
        time_id INTEGER NOT NULL,
        measure_type_id INTEGER NOT NULL,
        gender TEXT NOT NULL, -- Male, Female, Both
        incident_cases REAL,
        crude_rate REAL,
        age_adjusted_rate REAL,
        cumulative_risk_1_in_n INTEGER,
        deaths REAL,
        mortality_crude_rate REAL,
        FOREIGN KEY (site_id) REFERENCES Dim_Cancer_Site(site_id),
        FOREIGN KEY (time_id) REFERENCES Dim_Time(time_id),
        FOREIGN KEY (measure_type_id) REFERENCES Dim_Measure_Type(measure_type_id)
    );

    -- Fact 3: State-wise Annual Mortality & Incidence (Dataset 5.docx)
    CREATE TABLE Fact_State_Mortality_Surveillance (
        fact_mort_id INTEGER PRIMARY KEY AUTOINCREMENT,
        geo_id INTEGER NOT NULL,
        time_id INTEGER NOT NULL,
        measure_type_id INTEGER NOT NULL,
        reported_incidence REAL NOT NULL,
        reported_deaths REAL NOT NULL,
        mortality_incidence_ratio REAL,
        cancer_type_note TEXT DEFAULT 'UNIDENTIFIED - Dataset 5.docx does not state cancer type; scale (~77K national) suggests single-site (possibly ovarian); NOT all-cancer total',
        FOREIGN KEY (geo_id) REFERENCES Dim_Geography(geo_id),
        FOREIGN KEY (time_id) REFERENCES Dim_Time(time_id),
        FOREIGN KEY (measure_type_id) REFERENCES Dim_Measure_Type(measure_type_id)
    );

    -- Fact 4: Empirical PBCR Registry Surveillance
    CREATE TABLE Fact_PBCR_Surveillance (
        fact_pbcr_id INTEGER PRIMARY KEY AUTOINCREMENT,
        registry_id INTEGER NOT NULL,
        time_period TEXT NOT NULL,
        time_period_note TEXT DEFAULT 'Multi-year avg period as TEXT; not FK to Dim_Time - see PBCR publication',
        gender TEXT NOT NULL,
        annual_cases REAL,
        crude_rate REAL,
        age_adjusted_rate REAL,
        cum_risk_1_in_n INTEGER,
        tobacco_share_pct REAL,
        FOREIGN KEY (registry_id) REFERENCES Dim_Registry(registry_id)
    );

    -- Fact 5: Clinical Extent & Staging at Diagnosis (58 HBCRs)
    CREATE TABLE Fact_Clinical_Stage (
        fact_stage_id INTEGER PRIMARY KEY AUTOINCREMENT,
        source_name TEXT NOT NULL,
        source_measure_code TEXT DEFAULT 'JCO_2020_HBCR',
        cancer_site TEXT NOT NULL,
        gender TEXT NOT NULL,
        localized_pct REAL,
        locoregional_pct REAL,
        distant_metastasis_pct REAL,
        unknown_extent_pct REAL,
        advanced_stage_pct REAL
    );

    -- Fact 6: Registry Longitudinal Trends (APC %)
    CREATE TABLE Fact_Cancer_Trends_APC (
        fact_trend_id INTEGER PRIMARY KEY AUTOINCREMENT,
        cancer_site TEXT NOT NULL,
        gender TEXT NOT NULL,
        registry_name TEXT NOT NULL,
        region TEXT NOT NULL,
        apc_percent REAL,
        is_statistically_significant INTEGER,
        direction TEXT,
        source_measure_code TEXT DEFAULT 'JCO_2020_APC'
    );

    -- Fact 7: National Early Detection Screening (NP-NCD)
    CREATE TABLE Fact_Screening_Programs (
        screening_id INTEGER PRIMARY KEY AUTOINCREMENT,
        cancer_type TEXT NOT NULL,
        persons_screened INTEGER NOT NULL,
        cases_diagnosed INTEGER NOT NULL,
        cases_on_treatment INTEGER NOT NULL,
        detection_rate_per_100k REAL,
        treatment_retention_rate_pct REAL
    );

    -- Fact 8: National Healthcare Infrastructure & Approvals (Dataset 6.docx Table 1)
    CREATE TABLE Fact_Healthcare_Infrastructure_National (
        infra_id INTEGER PRIMARY KEY AUTOINCREMENT,
        indicator_name TEXT UNIQUE NOT NULL,
        facility_count INTEGER NOT NULL,
        reporting_source TEXT NOT NULL
    );

    -- Fact 9: State Health Infrastructure (Maharashtra Deepdive - fact sheet)
    CREATE TABLE Fact_Maharashtra_Infrastructure (
        mh_infra_id INTEGER PRIMARY KEY AUTOINCREMENT,
        category TEXT NOT NULL,
        facility_name TEXT NOT NULL,
        facility_count_or_type TEXT NOT NULL,
        source_note TEXT
    );

    -- Star Schema View 1: State Incidence & Projections
    CREATE VIEW vw_star_state_incidence AS
    SELECT 
        f.fact_state_id,
        g.state_name,
        g.state_code,
        g.region_zone,
        t.year,
        t.epoch_category,
        m.measure_code,
        m.measure_name,
        m.is_modeled,
        f.incident_cases,
        f.national_share_pct,
        f.cagr_observed,
        f.cagr_projected
    FROM Fact_State_Surveillance f
    JOIN Dim_Geography g ON f.geo_id = g.geo_id
    JOIN Dim_Time t ON f.time_id = t.time_id
    JOIN Dim_Measure_Type m ON f.measure_type_id = m.measure_type_id;

    -- Star Schema View 2: Cancer Sites & System Hierarchy
    CREATE VIEW vw_star_cancer_sites AS
    SELECT 
        s.site_id,
        s.site_name,
        s.icd10_code,
        s.level_type,
        s.parent_group,
        s.is_tobacco_associated,
        o.system_name AS organ_system_name
    FROM Dim_Cancer_Site s
    LEFT JOIN Dim_Organ_System o ON s.organ_system_id = o.organ_system_id;

    -- Star Schema View 3: State Mortality Ratios (Dataset 5.docx)
    CREATE VIEW vw_state_mortality_ratios AS
    SELECT 
        m.fact_mort_id,
        g.state_name,
        g.state_code,
        g.region_zone,
        t.year,
        m.reported_incidence,
        m.reported_deaths,
        m.mortality_incidence_ratio,
        m.cancer_type_note,
        mt.measure_code,
        mt.measure_name
    FROM Fact_State_Mortality_Surveillance m
    JOIN Dim_Geography g ON m.geo_id = g.geo_id
    JOIN Dim_Time t ON m.time_id = t.time_id
    JOIN Dim_Measure_Type mt ON m.measure_type_id = mt.measure_type_id;
    """)
    conn.commit()
    logging.info("Schema creation completed successfully.")

def populate_dimensions(conn, wb_updated):
    """Populate Dim_Geography, Dim_Time, Dim_Measure_Type, and Dim_Organ_System."""
    cursor = conn.cursor()
    logging.info("Populating Conformed Dimensions...")

    # 1. Populate Dim_Time (time_id, year, period_label, epoch_category)
    # Corrected epoch labels: 2022 is Extrapolated NCRP; 2024 is Benchmark; 2019-2023 are Extrapolated
    times = [
        (2015, 2015, "2015 Historical", "Historical"),
        (2019, 2019, "2019 NCRP Extrapolation", "Extrapolated"),
        (2020, 2020, "2020 NCRP Extrapolation", "Extrapolated"),
        (2021, 2021, "2021 Estimated (PARLIAMENT)", "Extrapolated"),
        (2022, 2022, "2022 Extrapolated NCRP", "Extrapolated"),
        (2023, 2023, "2023 Estimated (PARLIAMENT)", "Extrapolated"),
        (2024, 2024, "2024 GLOBOCAN Benchmark", "Benchmark"),
        (2025, 2025, "2025 Projection", "Projected"),
        (2026, 2026, "2026 Forecast", "Projected"),
        (2030, 2030, "2030 Forecast", "Projected"),
        (2045, 2045, "2045 Forecast", "Projected")
    ]
    cursor.executemany("INSERT INTO Dim_Time (time_id, year, period_label, epoch_category) VALUES (?, ?, ?, ?)", times)

    # 2. Populate Dim_Measure_Type (Clear, documented measure classification)
    measures = [
        ("PBCR_EMPIRICAL", "Empirical PBCR Registry Registrations", "Direct registry data pooled from 28 Population-Based Cancer Registries", 0),
        ("ICMR_NCRP_ESTIMATE", "ICMR-NCRP National Extrapolated Estimate", "State/National estimates applying 2012-2016 PBCR rates to Census population", 0),
        ("KUMAR_EXP_2026", "Moderated Exponential Model Forecast", "Kumar et al. 2026 (IJCMPH) econometric forecasting model", 1),
        ("GLOBOCAN_2024", "WHO GLOBOCAN 2024 Benchmark", "IARC / World Health Organization standard international cancer estimates", 0),
        ("PARLIAMENT_MOHFW", "Parliamentary Surveillance Response (MOHFW)", "Ministry of Health & Family Welfare state-level surveillance responses", 0)
    ]
    cursor.executemany("INSERT INTO Dim_Measure_Type (measure_code, measure_name, methodology_source, is_modeled) VALUES (?, ?, ?, ?)", measures)

    # 3. Populate Dim_Geography from State-wise Incidence
    ws_state = wb_updated["State-wise Incidence"]
    geo_records = []
    seen_states = set()
    for row in list(ws_state.iter_rows(values_only=True))[1:]:
        sname = row[0]
        if not sname or "all india" in str(sname).lower() or "source" in str(sname).lower() or "added" in str(sname).lower() or "caveat" in str(sname).lower():
            continue
        sname = str(sname).strip()
        if sname in seen_states:
            continue
        seen_states.add(sname)
        code, region = STATE_CODE_MAP.get(sname, ("XX", row[8] if len(row) > 8 and row[8] else "Other"))
        geo_records.append((sname, code, region))

    cursor.executemany("INSERT INTO Dim_Geography (state_name, state_code, region_zone) VALUES (?, ?, ?)", geo_records)
    logging.info(f"Loaded {len(geo_records)} distinct states into Dim_Geography.")

    # 4. Populate Dim_Organ_System
    organ_systems = [
        ("Digestive System",),
        ("Respiratory & Intrathoracic",),
        ("Oral Cavity & Pharynx",),
        ("Female Genital Organs",),
        ("Breast",),
        ("Male Genital Organs",),
        ("Urinary Tract",),
        ("Lymphoid & Haematopoietic",),
        ("Head & Neck",),
        ("Endocrine & Other",),
        ("All Sites Combined",)
    ]
    cursor.executemany("INSERT INTO Dim_Organ_System (system_name) VALUES (?)", organ_systems)
    conn.commit()

def populate_cancer_sites(conn, wb_updated, wb9):
    """
    Populate Dim_Cancer_Site.
    Disambiguate pharynx and other sub-sites so Hypopharynx, Oropharynx,
    Nasopharynx, Gallbladder, etc. remain distinct without string collision.
    """
    cursor = conn.cursor()
    logging.info("Populating Dim_Cancer_Site with exact disambiguation...")

    cursor.execute("SELECT system_name, organ_system_id FROM Dim_Organ_System")
    sys_map = dict(cursor.fetchall())

    ws_sites = wb_updated["India Sites 2022"]
    site_records = []
    seen = set()

    tobacco_sites = {
        "Lip, oral cavity", "Mouth", "Tongue", "Pharynx", "Other oral cavity",
        "Oesophagus", "Larynx", "Lung", "Lung and bronchus", "Urinary bladder",
        "Hypopharynx", "Oropharynx", "Nasopharynx"
    }

    def get_organ_system(site_name, parent):
        s_lower = site_name.lower()
        p_lower = str(parent).lower() if parent else ""
        if "all site" in s_lower or "all cancer" in s_lower:
            return sys_map["All Sites Combined"]
        if "oral" in s_lower or "mouth" in s_lower or "tongue" in s_lower or "pharynx" in s_lower or "oral" in p_lower:
            return sys_map["Oral Cavity & Pharynx"]
        if "digestive" in s_lower or "stomach" in s_lower or "colon" in s_lower or "rectum" in s_lower or "liver" in s_lower or "gallbladder" in s_lower or "pancreas" in s_lower or "colorectum" in s_lower or "digestive" in p_lower:
            return sys_map["Digestive System"]
        if "lung" in s_lower or "larynx" in s_lower or "trachea" in s_lower or "respiratory" in p_lower:
            return sys_map["Respiratory & Intrathoracic"]
        if "breast" in s_lower:
            return sys_map["Breast"]
        if "cervix" in s_lower or "ovary" in s_lower or "uterus" in s_lower or "vagina" in s_lower or "vulva" in s_lower:
            return sys_map["Female Genital Organs"]
        if "prostate" in s_lower or "testis" in s_lower or "penis" in s_lower:
            return sys_map["Male Genital Organs"]
        if "kidney" in s_lower or "bladder" in s_lower:
            return sys_map["Urinary Tract"]
        if "leukaemia" in s_lower or "lymphoma" in s_lower or "myeloma" in s_lower:
            return sys_map["Lymphoid & Haematopoietic"]
        return sys_map["Endocrine & Other"]

    # 1. Ingest IJMR 2022 sites
    for row in list(ws_sites.iter_rows(values_only=True))[1:]:
        site_raw = row[0]
        if not site_raw or "source" in str(site_raw).lower():
            continue
        site_name = str(site_raw).strip()
        level = str(row[1]).strip() if row[1] else "Site"
        parent = str(row[2]).strip() if row[2] else None
        
        icd = None
        match = re.search(r"\(([C\d\-, ]+)\)", site_name)
        if match:
            icd = match.group(1)
            site_clean = site_name[:match.start()].strip()
        else:
            site_clean = site_name

        if site_clean in seen:
            continue
        seen.add(site_clean)

        is_tob = 1 if any(t.lower() in site_clean.lower() for t in tobacco_sites) else 0
        sys_id = get_organ_system(site_clean, parent)
        site_records.append((site_clean, icd, level, parent, is_tob, sys_id))

    # 2. Add distinct GLOBOCAN 2024 sites that are not in IJMR
    if wb9 and "GLOBOCAN 2024 by Site" in wb9.sheetnames:
        ws_g = wb9["GLOBOCAN 2024 by Site"]
        for row in list(ws_g.iter_rows(values_only=True))[1:]:
            g_site = row[0]
            if not g_site or "source" in str(g_site).lower() or str(g_site).strip() == "None":
                continue
            g_name = str(g_site).strip()
            if g_name in seen:
                continue
            
            # Map parent and specific ICD if standard
            icd = None
            parent = None
            if g_name in ("Hypopharynx", "Oropharynx", "Nasopharynx"):
                parent = "Oral cavity and pharynx"
            elif g_name == "Gallbladder":
                parent = "Digestive system"
            elif g_name == "Colorectum":
                parent = "Digestive system"
            
            seen.add(g_name)
            is_tob = 1 if any(t.lower() in g_name.lower() for t in tobacco_sites) else 0
            sys_id = get_organ_system(g_name, parent)
            site_records.append((g_name, icd, "Site", parent, is_tob, sys_id))

    cursor.executemany("""
        INSERT INTO Dim_Cancer_Site (site_name, icd10_code, level_type, parent_group, is_tobacco_associated, organ_system_id)
        VALUES (?, ?, ?, ?, ?, ?)
    """, site_records)
    conn.commit()
    logging.info(f"Loaded {len(site_records)} distinct cancer sites into Dim_Cancer_Site.")

def load_fact_state_surveillance(conn, wb_updated):
    """
    Load Fact_State_Surveillance.
    2019-2022 state figures classified as ICMR_NCRP_ESTIMATE,
    and 2026-2045 as KUMAR_EXP_2026.
    """
    cursor = conn.cursor()
    logging.info("Loading Fact_State_Surveillance with corrected measure classification...")

    cursor.execute("SELECT state_name, geo_id FROM Dim_Geography")
    geo_map = dict(cursor.fetchall())
    
    cursor.execute("SELECT measure_code, measure_type_id FROM Dim_Measure_Type")
    m_map = dict(cursor.fetchall())

    icmr_est_type = m_map["ICMR_NCRP_ESTIMATE"]
    proj_type = m_map["KUMAR_EXP_2026"]

    ws = wb_updated["State-wise Incidence"]
    rows = list(ws.iter_rows(values_only=True))[1:]
    
    records = []
    for r in rows:
        sname = str(r[0]).strip() if r[0] else ""
        if sname not in geo_map:
            continue
        gid = geo_map[sname]
        cagr_obs = clean_num(r[10])
        cagr_proj = clean_num(r[11])
        share_22 = clean_num(r[13])

        # ICMR Extrapolated Estimates: 2019 (col 1), 2020 (col 2), 2021 (col 3), 2022 (col 4)
        for yr, col_idx in [(2019, 1), (2020, 2), (2021, 3), (2022, 4)]:
            cases = clean_num(r[col_idx])
            if cases is not None:
                records.append((gid, yr, icmr_est_type, cases, cagr_obs, None, share_22 if yr == 2022 else None))

        # Model forecast years: 2026 (col 5), 2030 (col 6), 2045 (col 7)
        for yr, col_idx in [(2026, 5), (2030, 6), (2045, 7)]:
            cases = clean_num(r[col_idx])
            if cases is not None:
                records.append((gid, yr, proj_type, cases, None, cagr_proj, None))

    cursor.executemany("""
        INSERT INTO Fact_State_Surveillance (geo_id, time_id, measure_type_id, incident_cases, cagr_observed, cagr_projected, national_share_pct)
        VALUES (?, ?, ?, ?, ?, ?, ?)
    """, records)
    conn.commit()
    logging.info(f"Loaded {len(records)} state surveillance records.")

def load_fact_cancer_incidence(conn, wb_updated, wb9):
    """
    Load Fact_Cancer_Incidence.
    Exact disambiguation mapping for GLOBOCAN 2024 by Site,
    preventing duplicate rows for pharynx, gallbladder, etc.
    """
    cursor = conn.cursor()
    logging.info("Loading Fact_Cancer_Incidence (Site-wise national burden)...")

    cursor.execute("SELECT site_name, site_id FROM Dim_Cancer_Site")
    site_map = {name.lower(): sid for name, sid in cursor.fetchall()}

    cursor.execute("SELECT measure_code, measure_type_id FROM Dim_Measure_Type")
    m_map = dict(cursor.fetchall())
    icmr_type = m_map["ICMR_NCRP_ESTIMATE"]
    globo_type = m_map["GLOBOCAN_2024"]

    records = []

    # 1. From final_excel_updated: 'India Sites 2022'
    ws_sites = wb_updated["India Sites 2022"]
    for r in list(ws_sites.iter_rows(values_only=True))[1:]:
        site_raw = r[0]
        if not site_raw or "source" in str(site_raw).lower():
            continue
        site_str = str(site_raw).strip()
        match = re.search(r"\(([C\d\-, ]+)\)", site_str)
        clean_name = site_str[:match.start()].strip().lower() if match else site_str.lower()
        sid = site_map.get(clean_name)
        if not sid:
            continue

        # Male 2022 (cols 3, 4, 5, 6)
        m_cases = clean_num(r[3])
        m_cr = clean_num(r[4])
        m_aar = clean_num(r[5])
        m_risk = clean_num(r[6])
        if m_cases is not None:
            records.append((sid, 2022, icmr_type, "Male", m_cases, m_cr, m_aar, m_risk, None, None))

        # Female 2022 (cols 7, 8, 9, 10)
        f_cases = clean_num(r[7])
        f_cr = clean_num(r[8])
        f_aar = clean_num(r[9])
        f_risk = clean_num(r[10])
        if f_cases is not None:
            records.append((sid, 2022, icmr_type, "Female", f_cases, f_cr, f_aar, f_risk, None, None))

        # Both sexes 2022 (cols 11, 12)
        b_cases = clean_num(r[11])
        b_cr = clean_num(r[12])
        if b_cases is not None:
            records.append((sid, 2022, icmr_type, "Both", b_cases, b_cr, None, None, None, None))

    # 2. From Dataset 9: 'GLOBOCAN 2024 by Site' with EXACT and dictionary-based resolution
    globo_alias = {
        "cervix uteri": "uterine cervix",
        "corpus uteri": "uterine corpus",
        "lung": "lung and bronchus",
        "liver": "liver and intrahepatic bile duct",
        "lip, oral cavity": "oral cavity and pharynx",
        "brain cns": "brain and other nervous system",
        "bladder": "urinary bladder",
        "nhl": "non-hodgkin lymphoma",
        "kidney": "kidney and renal pelvis",
        "all cancers": "all sites"
    }

    if wb9 and "GLOBOCAN 2024 by Site" in wb9.sheetnames:
        ws_g = wb9["GLOBOCAN 2024 by Site"]
        for r in list(ws_g.iter_rows(values_only=True))[1:]:
            c_raw = r[0]
            if not c_raw or "source" in str(c_raw).lower() or str(c_raw).strip() == "None":
                continue
            c_name = str(c_raw).strip()
            c_lower = c_name.lower()
            
            # Step A: Exact match in Dim_Cancer_Site
            sid = site_map.get(c_lower)
            # Step B: Explicit clinical alias
            if not sid and c_lower in globo_alias:
                sid = site_map.get(globo_alias[c_lower])

            if not sid:
                logging.warning(f"GLOBOCAN site could not be mapped: {c_name}")
                continue

            cases = clean_num(r[1])
            deaths = clean_num(r[5])
            if cases is not None:
                records.append((sid, 2024, globo_type, "Both", cases, None, None, None, deaths, None))

    cursor.executemany("""
        INSERT INTO Fact_Cancer_Incidence (site_id, time_id, measure_type_id, gender, incident_cases, crude_rate, age_adjusted_rate, cumulative_risk_1_in_n, deaths, mortality_crude_rate)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, records)
    conn.commit()
    logging.info(f"Loaded {len(records)} site-level incidence records.")

def load_fact_state_mortality(conn, docx_path):
    """
    Ingest Dataset 5.docx state-wise annual incidence and mortality (2021-2025).
    Tagged with PARLIAMENT_MOHFW measure code and explicit cancer_type_note documenting uncertainty.
    """
    cursor = conn.cursor()
    logging.info("Loading Fact_State_Mortality_Surveillance from Dataset 5.docx...")

    cursor.execute("SELECT state_name, geo_id FROM Dim_Geography")
    geo_map = dict(cursor.fetchall())

    cursor.execute("SELECT measure_type_id FROM Dim_Measure_Type WHERE measure_code = 'PARLIAMENT_MOHFW'")
    m_type_id = cursor.fetchone()[0]

    # Mapping for states with historical or alternate spellings
    alias_lookup = {
        "uttarakhand": "uttaranchal",
        "odisha": "orissa",
        "puducherry": "pondicherry",
        "andaman & nicobar islands": "andaman and nicobar islands",
        "jammu & kashmir": "jammu and kashmir",
        "daman & diu": "daman",
        "dadra & nagar haveli": "dadra and nagar haveli"
    }

    doc = docx.Document(docx_path)
    t = doc.tables[0]
    
    cancer_note = "UNIDENTIFIED - Dataset 5.docx does not state cancer type; scale (~77K national) suggests single-site (possibly ovarian); NOT all-cancer total"

    # Headers: State/UT, Inc 2021..2025 (cols 1-5), Deaths 2021..2025 (cols 6-10)
    records = []
    for row in t.rows[1:]:
        sname_raw = row.cells[0].text.strip()
        if not sname_raw or "total" in sname_raw.lower():
            continue
        
        # Match geography
        sname_clean = sname_raw.lower()
        if sname_clean in alias_lookup:
            sname_clean = alias_lookup[sname_clean]
        
        gid = None
        for k, v in geo_map.items():
            if k.lower() == sname_clean or k.lower().replace("&", "and") == sname_clean.replace("&", "and"):
                gid = v
                break
        if not gid:
            logging.warning(f"Could not map Dataset 5 state: {sname_raw}")
            continue

        for yr_idx, yr in enumerate([2021, 2022, 2023, 2024, 2025]):
            inc_val = clean_num(row.cells[1 + yr_idx].text)
            death_val = clean_num(row.cells[6 + yr_idx].text)

            if inc_val is not None and death_val is not None:
                mir = round(death_val / inc_val, 4) if inc_val > 0 else None
                records.append((gid, yr, m_type_id, inc_val, death_val, mir, cancer_note))

    cursor.executemany("""
        INSERT INTO Fact_State_Mortality_Surveillance (geo_id, time_id, measure_type_id, reported_incidence, reported_deaths, mortality_incidence_ratio, cancer_type_note)
        VALUES (?, ?, ?, ?, ?, ?, ?)
    """, records)
    conn.commit()
    logging.info(f"Loaded {len(records)} state mortality surveillance records.")

def load_fact_pbcr_surveillance(conn, wb_updated):
    """Load empirical PBCR registry records and tobacco attributions."""
    cursor = conn.cursor()
    logging.info("Loading Fact_PBCR_Surveillance...")

    cursor.execute("SELECT measure_type_id FROM Dim_Measure_Type WHERE measure_code = 'PBCR_EMPIRICAL'")
    emp_type = cursor.fetchone()[0]

    ws_pbcr = wb_updated["PBCR Incidence Table"]
    rows = list(ws_pbcr.iter_rows(values_only=True))[1:]
    
    for r in rows:
        reg_name = r[0]
        if not reg_name or "source" in str(reg_name).lower():
            continue
        reg_name = str(reg_name).strip()
        region = str(r[2]).strip() if r[2] else "Other"
        cursor.execute("INSERT OR IGNORE INTO Dim_Registry (registry_name, region) VALUES (?, ?)", (reg_name, region))
    conn.commit()

    cursor.execute("SELECT registry_name, registry_id FROM Dim_Registry")
    registry_map = dict(cursor.fetchall())

    tobacco_map = {}
    if "Tobacco-related Cancers" in wb_updated.sheetnames:
        ws_tob = wb_updated["Tobacco-related Cancers"]
        for r in list(ws_tob.iter_rows(values_only=True))[1:]:
            rname = str(r[0]).strip() if r[0] else ""
            if rname:
                m_tob = clean_num(r[2])
                f_tob = clean_num(r[5])
                tobacco_map[rname] = {"Male": m_tob, "Female": f_tob}

    records = []
    period_note = "Multi-year avg period as TEXT; not FK to Dim_Time - see PBCR publication"
    for r in rows:
        reg_name = str(r[0]).strip() if r[0] else ""
        if not reg_name or "source" in str(reg_name).lower():
            continue
        if reg_name not in registry_map:
            continue
        rid = registry_map[reg_name]
        period = str(r[4]).strip() if r[4] else "2012-2016"

        tob_m = tobacco_map.get(reg_name, {}).get("Male")
        tob_f = tobacco_map.get(reg_name, {}).get("Female")

        # Male (cols 5, 6, 7, 8)
        m_cases = clean_num(r[5])
        m_cr = clean_num(r[6])
        m_aar = clean_num(r[7])
        m_risk = clean_num(r[8])
        if m_cases is not None:
            records.append((rid, period, period_note, "Male", m_cases, m_cr, m_aar, m_risk, tob_m))

        # Female (cols 10, 11, 12, 13)
        f_cases = clean_num(r[10])
        f_cr = clean_num(r[11])
        f_aar = clean_num(r[12])
        f_risk = clean_num(r[13])
        if f_cases is not None:
            records.append((rid, period, period_note, "Female", f_cases, f_cr, f_aar, f_risk, tob_f))

    cursor.executemany("""
        INSERT INTO Fact_PBCR_Surveillance (registry_id, time_period, time_period_note, gender, annual_cases, crude_rate, age_adjusted_rate, cum_risk_1_in_n, tobacco_share_pct)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, records)
    conn.commit()
    logging.info(f"Loaded {len(records)} empirical PBCR surveillance records.")

def load_fact_clinical_stage(conn, wb_updated):
    """Load clinical staging distribution from Hospital-Based Registries (58 HBCRs)."""
    cursor = conn.cursor()
    logging.info("Loading Fact_Clinical_Stage...")

    ws = wb_updated["Clinical Stage"]
    records = []
    for r in list(ws.iter_rows(values_only=True))[1:]:
        src = r[0]
        if not src or "sources:" in str(src).lower():
            continue
        csite = str(r[2]).strip()
        sex = str(r[3]).strip()
        loc = clean_num(r[4])
        locoreg = clean_num(r[5])
        dist = clean_num(r[6])
        unk = clean_num(r[7])
        adv = clean_num(r[9])

        records.append((str(src).strip(), "JCO_2020_HBCR", csite, sex, loc, locoreg, dist, unk, adv))

    cursor.executemany("""
        INSERT INTO Fact_Clinical_Stage (source_name, source_measure_code, cancer_site, gender, localized_pct, locoregional_pct, distant_metastasis_pct, unknown_extent_pct, advanced_stage_pct)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, records)
    conn.commit()
    logging.info(f"Loaded {len(records)} clinical stage records.")

def load_fact_trends_apc(conn, wb_updated):
    """Load longitudinal Annual Percentage Change (APC) trends, skipping stray metadata rows."""
    cursor = conn.cursor()
    logging.info("Loading Fact_Cancer_Trends_APC...")

    ws = wb_updated["Cancer Trends (APC)"]
    records = []
    for r in list(ws.iter_rows(values_only=True))[1:]:
        site = r[0]
        if not site:
            continue
        site_str = str(site).strip()

        # Skip stray footnote/metadata rows
        if site_str.lower().startswith("source:"):
            logging.warning(f"Skipping footnote row in Cancer Trends APC: {site_str[:80]}")
            continue

        sex = str(r[1]).strip()
        reg = str(r[2]).strip()
        region = str(r[3]).strip() if r[3] else "Other"
        apc = clean_num(r[4])
        sig = 1 if str(r[5]).strip().lower() == "yes" else 0
        direction = str(r[6]).strip() if r[6] else ("Increase" if apc and apc > 0 else "Decrease")

        records.append((site_str, sex, reg, region, apc, sig, direction, "JCO_2020_APC"))

    cursor.executemany("""
        INSERT INTO Fact_Cancer_Trends_APC (cancer_site, gender, registry_name, region, apc_percent, is_statistically_significant, direction, source_measure_code)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    """, records)
    conn.commit()
    logging.info(f"Loaded {len(records)} APC longitudinal trend records.")

def load_fact_screening(conn, docx_path):
    """Load national NP-NCD early detection screening data from Dataset 6.docx."""
    cursor = conn.cursor()
    logging.info("Loading Fact_Screening_Programs from Dataset 6.docx...")

    doc = docx.Document(docx_path)
    t2 = doc.tables[1]  # Screening Table
    records = []
    for row in t2.rows[1:]:
        c_type = row.cells[0].text.strip()
        screened = clean_num(row.cells[1].text)
        diag = clean_num(row.cells[2].text)
        tx = clean_num(row.cells[3].text)

        det_rate = round((diag / screened) * 100000, 2) if screened and diag else None
        tx_retention = round((tx / diag) * 100, 2) if diag and tx else None

        records.append((c_type, screened, diag, tx, det_rate, tx_retention))

    cursor.executemany("""
        INSERT INTO Fact_Screening_Programs (cancer_type, persons_screened, cases_diagnosed, cases_on_treatment, detection_rate_per_100k, treatment_retention_rate_pct)
        VALUES (?, ?, ?, ?, ?, ?)
    """, records)
    conn.commit()
    logging.info(f"Loaded {len(records)} national screening program records.")

def load_fact_healthcare_infrastructure(conn, docx_path, wb_updated):
    """
    Ingest verified healthcare infrastructure data from Dataset 6.docx (Table 1)
    and Maharashtra Health Infrastructure (Fact sheet).
    """
    cursor = conn.cursor()
    logging.info("Loading Healthcare Infrastructure facts (National and State)...")

    # 1. National Approvals (Dataset 6.docx Table 1)
    if os.path.exists(docx_path):
        doc = docx.Document(docx_path)
        t1 = doc.tables[0]
        nat_records = []
        for r in t1.rows[1:]:
            ind = r.cells[0].text.strip()
            cnt = clean_num(r.cells[1].text)
            if ind and cnt is not None:
                nat_records.append((ind, cnt, "MOHFW Parliamentary Response (Feb 2026)"))
        cursor.executemany("""
            INSERT OR REPLACE INTO Fact_Healthcare_Infrastructure_National (indicator_name, facility_count, reporting_source)
            VALUES (?, ?, ?)
        """, nat_records)
        logging.info(f"Loaded {len(nat_records)} national infrastructure indicators.")

    # 2. Maharashtra Infrastructure (final_excel_updated.xlsx)
    if "MH Health Infrastructure" in wb_updated.sheetnames:
        ws_mhi = wb_updated["MH Health Infrastructure"]
        mh_records = []
        for r in list(ws_mhi.iter_rows(values_only=True))[1:]:
            cat = str(r[0]).strip() if r[0] else ""
            item = str(r[1]).strip() if r[1] else ""
            cnt_detail = str(r[2]).strip() if r[2] else ""
            src = str(r[4]).strip() if len(r) > 4 and r[4] else ""
            if cat and item:
                mh_records.append((cat, item, cnt_detail, src))
        cursor.executemany("""
            INSERT INTO Fact_Maharashtra_Infrastructure (category, facility_name, facility_count_or_type, source_note)
            VALUES (?, ?, ?, ?)
        """, mh_records)
        logging.info(f"Loaded {len(mh_records)} Maharashtra health infrastructure records.")

    conn.commit()

def run_integrity_checks(conn):
    """Run automated validation checks and return results dict."""
    cursor = conn.cursor()
    logging.info("Running Data Integrity and Validation Suite...")
    results = {}

    # Check 1: Row counts per table
    tables = [
        "Dim_Geography", "Dim_Cancer_Site", "Dim_Organ_System", "Dim_Time", "Dim_Measure_Type", "Dim_Registry",
        "Fact_State_Surveillance", "Fact_Cancer_Incidence", "Fact_State_Mortality_Surveillance",
        "Fact_PBCR_Surveillance", "Fact_Clinical_Stage", "Fact_Cancer_Trends_APC",
        "Fact_Screening_Programs", "Fact_Healthcare_Infrastructure_National", "Fact_Maharashtra_Infrastructure"
    ]
    counts = {}
    for t in tables:
        cursor.execute(f"SELECT COUNT(*) FROM {t}")
        counts[t] = cursor.fetchone()[0]
    results["row_counts"] = counts
    logging.info(f"Row counts: {json.dumps(counts, indent=2)}")

    # Check 2: Sum of 37 state cases for 2022 vs reported national estimate (1,461,427)
    cursor.execute("""
        SELECT SUM(incident_cases) 
        FROM Fact_State_Surveillance 
        WHERE time_id = 2022 AND measure_type_id = (SELECT measure_type_id FROM Dim_Measure_Type WHERE measure_code = 'ICMR_NCRP_ESTIMATE')
    """)
    sum_2022 = cursor.fetchone()[0]
    expected_2022 = 1461427
    diff_2022 = abs(sum_2022 - expected_2022) if sum_2022 else expected_2022
    pct_diff = (diff_2022 / expected_2022) * 100 if expected_2022 else 100
    results["national_sum_check_2022"] = {
        "calculated_sum": sum_2022,
        "expected_sum": expected_2022,
        "absolute_difference": diff_2022,
        "percentage_error": round(pct_diff, 4),
        "status": "PASS" if pct_diff < 0.01 else "FAIL"
    }
    logging.info(f"Sum validation: Calculated={sum_2022}, Expected={expected_2022}, Status={results['national_sum_check_2022']['status']}")

    # Check 3: Orphan Foreign Keys
    cursor.execute("SELECT COUNT(*) FROM Fact_State_Surveillance WHERE geo_id NOT IN (SELECT geo_id FROM Dim_Geography)")
    orphan_geo = cursor.fetchone()[0]
    cursor.execute("SELECT COUNT(*) FROM Fact_Cancer_Incidence WHERE site_id NOT IN (SELECT site_id FROM Dim_Cancer_Site)")
    orphan_site = cursor.fetchone()[0]
    cursor.execute("SELECT COUNT(*) FROM Fact_State_Mortality_Surveillance WHERE geo_id NOT IN (SELECT geo_id FROM Dim_Geography)")
    orphan_mort_geo = cursor.fetchone()[0]
    cursor.execute("SELECT COUNT(*) FROM Fact_State_Surveillance WHERE time_id NOT IN (SELECT time_id FROM Dim_Time)")
    orphan_time = cursor.fetchone()[0]
    cursor.execute("SELECT COUNT(*) FROM Fact_State_Surveillance WHERE measure_type_id NOT IN (SELECT measure_type_id FROM Dim_Measure_Type)")
    orphan_measure = cursor.fetchone()[0]
    total_orphans = orphan_geo + orphan_site + orphan_mort_geo + orphan_time + orphan_measure
    results["referential_integrity"] = {
        "orphan_state_surveillance_geos": orphan_geo,
        "orphan_cancer_incidence_sites": orphan_site,
        "orphan_state_mortality_geos": orphan_mort_geo,
        "orphan_state_surveillance_times": orphan_time,
        "orphan_state_surveillance_measures": orphan_measure,
        "status": "PASS" if total_orphans == 0 else "FAIL"
    }

    # Check 4: Duplicate logical records
    cursor.execute("""
        SELECT site_id, time_id, measure_type_id, gender, COUNT(*) 
        FROM Fact_Cancer_Incidence 
        GROUP BY site_id, time_id, measure_type_id, gender 
        HAVING COUNT(*) > 1
    """)
    dups_incidence = cursor.fetchall()
    results["duplicate_logical_records"] = {
        "fact_cancer_incidence_duplicates": len(dups_incidence),
        "status": "PASS" if len(dups_incidence) == 0 else "WARNING"
    }

    # Check 5: Domain boundary checks
    cursor.execute("SELECT COUNT(*) FROM Fact_Clinical_Stage WHERE advanced_stage_pct < 0 OR advanced_stage_pct > 100")
    invalid_stages = cursor.fetchone()[0]
    cursor.execute("SELECT COUNT(*) FROM Fact_State_Mortality_Surveillance WHERE mortality_incidence_ratio < 0 OR mortality_incidence_ratio > 1.5")
    invalid_mirs = cursor.fetchone()[0]
    results["domain_boundary_checks"] = {
        "invalid_clinical_stage_percentages": invalid_stages,
        "invalid_mortality_incidence_ratios": invalid_mirs,
        "status": "PASS" if invalid_stages == 0 and invalid_mirs == 0 else "FAIL"
    }

    # Check 6: Stray footnote rows in Fact_Cancer_Trends_APC
    cursor.execute("SELECT COUNT(*) FROM Fact_Cancer_Trends_APC WHERE LOWER(cancer_site) LIKE 'source:%'")
    stray_apc = cursor.fetchone()[0]
    results["stray_footnote_rows"] = {
        "stray_cancer_trends_rows": stray_apc,
        "status": "PASS" if stray_apc == 0 else "FAIL"
    }

    # Check 7: NULL in required fields
    cursor.execute("SELECT COUNT(*) FROM Fact_State_Surveillance WHERE incident_cases IS NULL")
    null_cases = cursor.fetchone()[0]
    cursor.execute("SELECT COUNT(*) FROM Fact_State_Mortality_Surveillance WHERE reported_incidence IS NULL OR reported_deaths IS NULL")
    null_mort = cursor.fetchone()[0]
    results["null_required_fields"] = {
        "fact_state_surveillance_null_cases": null_cases,
        "fact_state_mortality_nulls": null_mort,
        "status": "PASS" if null_cases == 0 and null_mort == 0 else "FAIL"
    }

    # Check 8: Negative values
    cursor.execute("SELECT COUNT(*) FROM Fact_State_Surveillance WHERE incident_cases < 0")
    neg_cases = cursor.fetchone()[0]
    cursor.execute("SELECT COUNT(*) FROM Fact_State_Mortality_Surveillance WHERE reported_incidence < 0 OR reported_deaths < 0")
    neg_mort = cursor.fetchone()[0]
    results["negative_values"] = {
        "negative_state_cases": neg_cases,
        "negative_mortality": neg_mort,
        "status": "PASS" if neg_cases == 0 and neg_mort == 0 else "FAIL"
    }

    return results

def main():
    logging.info("Starting Phase 2.1 Controlled ETL Pipeline...")

    if not os.path.exists(FINAL_EXCEL):
        logging.error(f"Missing required primary source: {FINAL_EXCEL}")
        sys.exit(1)

    wb_updated = openpyxl.load_workbook(FINAL_EXCEL, data_only=True)
    wb9 = openpyxl.load_workbook(DATASET_9, data_only=True) if os.path.exists(DATASET_9) else None

    # Ensure clean temp DB path
    if os.path.exists(TEMP_DB_PATH):
        try:
            os.remove(TEMP_DB_PATH)
            logging.info("Removed leftover temporary database.")
        except Exception as e:
            logging.warning(f"Could not remove existing temporary database: {e}")

    conn = sqlite3.connect(TEMP_DB_PATH)
    try:
        init_database(conn)
        populate_dimensions(conn, wb_updated)
        populate_cancer_sites(conn, wb_updated, wb9)
        load_fact_state_surveillance(conn, wb_updated)
        load_fact_cancer_incidence(conn, wb_updated, wb9)
        if os.path.exists(DATASET_5):
            load_fact_state_mortality(conn, DATASET_5)
        load_fact_pbcr_surveillance(conn, wb_updated)
        load_fact_clinical_stage(conn, wb_updated)
        load_fact_trends_apc(conn, wb_updated)
        if os.path.exists(DATASET_6):
            load_fact_screening(conn, DATASET_6)
            load_fact_healthcare_infrastructure(conn, DATASET_6, wb_updated)
        
        val_results = run_integrity_checks(conn)
        conn.close()

        # Check critical validations before promoting temp DB
        critical_passed = True
        critical_keys = [
            "referential_integrity",
            "stray_footnote_rows",
            "negative_values",
            "null_required_fields",
            "domain_boundary_checks"
        ]
        for key in critical_keys:
            if key in val_results and val_results[key].get("status") != "PASS":
                logging.error(f"Critical validation check failed: {key} -> {val_results[key]}")
                critical_passed = False

        if not critical_passed:
            logging.error("ETL Aborted: One or more critical validation checks failed. Active database untouched.")
            if os.path.exists(TEMP_DB_PATH):
                os.remove(TEMP_DB_PATH)
            sys.exit(1)

        # Atomic promotion: replace temp DB to active DB
        os.replace(TEMP_DB_PATH, DB_PATH)
        logging.info(f"Successfully promoted temporary database to active warehouse: {DB_PATH}")

        # Write validation results JSON
        with open(os.path.join("warehouse", "validation_results.json"), "w", encoding="utf-8") as f:
            json.dump(val_results, f, indent=2)

        summary_status = {k: v.get("status", "N/A") if isinstance(v, dict) else len(v) if isinstance(v, dict) else v for k, v in val_results.items()}
        logging.info(f"Phase 2.1 ETL Pipeline completed. Validation results: {json.dumps(summary_status, indent=2)}")

    except Exception as e:
        logging.error(f"ETL Execution Failed: {str(e)}", exc_info=True)
        try:
            conn.close()
        except Exception:
            pass
        if os.path.exists(TEMP_DB_PATH):
            try:
                os.remove(TEMP_DB_PATH)
                logging.info("Cleaned up temporary database after failure.")
            except Exception:
                pass
        sys.exit(1)

if __name__ == "__main__":
    main()
